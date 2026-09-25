# -*- coding: utf-8 -*-
"""
生成《鬼畜哈吉米 · 模块地图》工作簿。

内容分工：
  · 中文名 / 一句话作用 / 坑      → 人工撰写（data.py）
  · 谁在用它 / 它用了谁           → 从真实 import 边自动生成（module-graph.json），不手抄
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from data import CHAINS, MODULES  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
GRAPH = os.path.join(HERE, "..", ".workbuddy", "tmp", "module-graph.json")
OUT = os.path.join(HERE, "..", "鬼畜哈吉米模块地图.xlsx")

# ── 颜色（CSS #RRGGBB → openpyxl ARGB）────────────────────────────────────
def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEADER = xl_color("#4472C4")
XL_HEADER_FONT = xl_color("#FFFFFF")
XL_TITLE_FONT = xl_color("#1F3864")
XL_BAND = xl_color("#F2F6FC")
XL_TOTAL = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
XL_TRAP = xl_color("#C00000")

THIN = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_L = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TITLE_AL = Alignment(horizontal="left", vertical="center")

# ── 依赖图 ───────────────────────────────────────────────────────────────
with open(GRAPH, encoding="utf-8") as f:
    raw = json.load(f)
GRAPH_MAP = {r["rel"]: r for r in raw}
NAMES = {path: name for _, path, name, _, _ in MODULES}
LAYER_OF = {path: layer for layer, path, _, _, _ in MODULES}

# 只把「运行时真的会调」的边写进调用链：类型字典是纯类型导入，写进每一行都是噪音
TYPE_ONLY = {"src/model/types"}
MAX_SHOWN = 6


def label(path: str) -> str:
    return NAMES.get(path, path)


def chain_text(paths, skip_type_dict=True):
    """把一堆文件路径变成「中文名（路径）」串，超长则省略计数。"""
    seen, out = set(), []
    for p in paths:
        if p not in NAMES:  # 测试文件 / 未收录
            continue
        if skip_type_dict and p in TYPE_ONLY:
            continue
        if p in seen:
            continue
        seen.add(p)
        out.append(f"{label(p)}（{p}）")
    if not out:
        return "—" if not skip_type_dict else "（仅类型字典）"
    if len(out) <= MAX_SHOWN:
        return "\n".join(out)
    return "\n".join(out[:MAX_SHOWN]) + f"\n…等 {len(out)} 个"


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 1 · 分层总览
# ═══════════════════════════════════════════════════════════════════════════
LAYER_NOTE = {
    "启动与外壳": "网页的入口、外壳、全局类型和音频总开关。任何一页都跑在它们之上。",
    "页面": "四个主页面 + 导出对话框 + 一个纯计算的小模块。页面只做编排，逻辑都在下面的层里。",
    "组件·键位": "键盘本身：怎么摆、按键什么手感、演奏台和乐器包导入的界面。",
    "组件·卷帘": "钢琴卷帘：坐标换算、画布绘制、全部鼠标/键盘交互。",
    "组件·装配": "「给这个音符配哪段素材」的那套面板与状态机。",
    "组件·素材": "素材箱里的卡片，以及双击波形打开的切片器。",
    "组件·通用": "全站共用的界面零件（图标、顶栏、弹窗、旋钮、效果器可视化）。",
    "引擎·地基": "发声的骨架：音频总开关、效果总线、按键机、播放器、录制、导出。",
    "引擎·效果": "四个效果器的具体实现，以及它们的名册与出厂默认值。",
    "引擎·音高": "判断「这段素材原本是哪个音」：YIN、AI 模型、后台线程与故障处理。",
    "引擎·变调": "改音高/改时长的各种实现：自研 Rust 内核、第三方库、出口验收。",
    "引擎·MIDI": "只负责读 .mid 文件，不碰工程数据。",
    "数据模型": "整个工程的数据与规则：状态仓库、存档、键与音高的总规则、片段装配。",
    "React Hooks": "把「引擎 + 仓库」包成 React 能直接用的几个钩子。",
    "工具": "与业务无关的小工具：格式化、编号、搜索、键盘绑定查找。",
}

ws = wb.active
ws.title = "分层总览"

LAYER_ORDER = []
for layer, *_ in MODULES:
    if layer not in LAYER_ORDER:
        LAYER_ORDER.append(layer)
counts = {L: sum(1 for m in MODULES if m[0] == L) for L in LAYER_ORDER}
TOTAL = sum(counts.values())

TITLE_R, HEAD_R = 1, 3
DATA_R0 = HEAD_R + 1
DATA_R1 = DATA_R0 + len(LAYER_ORDER) - 1
TOTAL_R = DATA_R1 + 1

ws.cell(TITLE_R, 1, "鬼畜哈吉米 · 分层总览")
ws.merge_cells(start_row=TITLE_R, start_column=1, end_row=TITLE_R, end_column=4)
ws.cell(TITLE_R, 1).font = Font(size=14, bold=True, color=XL_TITLE_FONT)
ws.cell(TITLE_R, 1).alignment = TITLE_AL
ws.row_dimensions[TITLE_R].height = 26

for col, name in zip("ABCD", ["分层", "模块数", "占比", "这一层管什么"]):
    c = ws[f"{col}{HEAD_R}"]
    c.value = name
    c.font = Font(bold=True, color=XL_HEADER_FONT)
    c.fill = PatternFill("solid", fgColor=XL_HEADER)
    c.alignment = CENTER
    c.border = BORDER

# 模块数 / 占比写「算好的静态值」而不是公式：
#   · 这是一份代码快照，数字不会自己变，公式没有收益；
#   · 没有公式就没有「缓存值缺失」问题，微信/邮件附件、pandas(data_only=True)
#     之类的预览器一定能看到数字（本机没装 LibreOffice，公式引擎也装不上）。
for i, layer in enumerate(LAYER_ORDER):
    r = DATA_R0 + i
    ws.cell(r, 1, layer)
    ws.cell(r, 2, counts[layer])
    ws.cell(r, 3, counts[layer] / TOTAL)
    ws.cell(r, 4, LAYER_NOTE.get(layer, ""))
    ws.cell(r, 2).number_format = "0"
    ws.cell(r, 3).number_format = "0.0%"
    for col in range(1, 5):
        cell = ws.cell(r, col)
        cell.border = BORDER
        cell.alignment = WRAP_L if col == 4 else CENTER
        if i % 2 == 1:
            cell.fill = PatternFill("solid", fgColor=XL_BAND)

ws.cell(TOTAL_R, 1, "合计")
ws.cell(TOTAL_R, 2, TOTAL)
ws.cell(TOTAL_R, 3, 1.0)
ws.cell(TOTAL_R, 4, "全部产品代码模块（不含 *.test.ts 测试文件）")
ws.cell(TOTAL_R, 2).number_format = "0"
ws.cell(TOTAL_R, 3).number_format = "0.0%"
for col in range(1, 5):
    cell = ws.cell(TOTAL_R, col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor=XL_TOTAL)
    cell.border = BORDER
    cell.alignment = WRAP_L if col == 4 else CENTER

GUIDE_R = TOTAL_R + 2
guide = [
    ("怎么读这张表", ""),
    ("调用链的方向", "「谁在用它」= 上游，点它的地方；「它用了谁」= 下游，它去调的地方。两者都由真实 import 关系自动生成。"),
    ("为什么下游少了「全局类型字典」", "src/model/types.ts 是纯类型声明（编译后不存在），几乎每个文件都引用它。写进每一行只会淹没真正的调用关系，故省略。"),
    ("「坑」列", "红色文字 = 这个模块踩过坑或有一条硬规矩，改它之前建议先看说明。"),
    ("和记忆文档的关系", "更长的来龙去脉在项目 .workbuddy/memory/ 下的 REF-*.md 里；本表是索引与总览，不是替代。"),
]
for i, (k, v) in enumerate(guide):
    r = GUIDE_R + i
    ws.cell(r, 1, k).font = Font(bold=True, color=XL_TITLE_FONT if i == 0 else XL_HEADER)
    if i > 0:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.cell(r, 2, v).alignment = WRAP_L

ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 78

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 2 · 模块地图（主表）
# ═══════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("模块地图")
COLS = ["序号", "分层", "模块名", "文件", "这个模块干什么（大白话）",
        "谁在用它（上游）", "它用了谁（下游）", "改动时的坑"]
T2_R, H2_R = 1, 2
D2_R0 = 3

ws2.cell(T2_R, 1, "鬼畜哈吉米 · 模块地图（共 %d 个产品模块）" % TOTAL)
ws2.merge_cells(start_row=T2_R, start_column=1, end_row=T2_R, end_column=len(COLS))
ws2.cell(T2_R, 1).font = Font(size=14, bold=True, color=XL_TITLE_FONT)
ws2.cell(T2_R, 1).alignment = TITLE_AL
ws2.row_dimensions[T2_R].height = 26

for i, name in enumerate(COLS, start=1):
    c = ws2.cell(H2_R, i, name)
    c.font = Font(bold=True, color=XL_HEADER_FONT)
    c.fill = PatternFill("solid", fgColor=XL_HEADER)
    c.alignment = CENTER
    c.border = BORDER

for i, (layer, path, name, purpose, trap) in enumerate(MODULES):
    r = D2_R0 + i
    g = GRAPH_MAP.get(path, {})
    ws2.cell(r, 1, i)
    ws2.cell(r, 2, layer)
    ws2.cell(r, 3, name)
    ws2.cell(r, 4, path)
    ws2.cell(r, 5, purpose)
    ws2.cell(r, 6, chain_text(g.get("users", [])))
    ws2.cell(r, 7, chain_text(g.get("deps", [])))
    ws2.cell(r, 8, trap or "—")
    for col in range(1, len(COLS) + 1):
        cell = ws2.cell(r, col)
        cell.border = BORDER
        if col == 1:
            cell.alignment = CENTER
            cell.number_format = "0"
        elif col == 2:
            cell.alignment = CENTER
        elif col == 4:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.font = Font(name="Consolas", size=9)
        elif col == 3:
            cell.alignment = WRAP_L
            cell.font = Font(bold=True)
        elif col == 8:
            cell.alignment = WRAP_L
            if trap:
                cell.font = Font(color=XL_TRAP, size=10)
        else:
            cell.alignment = WRAP_L
        if i % 2 == 1 and col != 8:
            cell.fill = PatternFill("solid", fgColor=XL_BAND)

LAST2 = D2_R0 + len(MODULES) - 1
ws2.auto_filter.ref = f"A{H2_R}:H{LAST2}"
ws2.freeze_panes = f"C{D2_R0}"

for col, w in zip("ABCDEFGH", [6, 12, 18, 40, 46, 34, 34, 40]):
    ws2.column_dimensions[col].width = w

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 3 · 主调用链
# ═══════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("主调用链")
C3 = ["序号", "链路", "第几步", "这一步发生什么（大白话）", "经过哪个模块", "对应文件"]
T3_R, H3_R = 1, 2
D3_R0 = 3

ws3.cell(T3_R, 1, "鬼畜哈吉米 · 十三条主调用链（一件事从开始到结束，数据都经过谁）")
ws3.merge_cells(start_row=T3_R, start_column=1, end_row=T3_R, end_column=len(C3))
ws3.cell(T3_R, 1).font = Font(size=14, bold=True, color=XL_TITLE_FONT)
ws3.cell(T3_R, 1).alignment = TITLE_AL
ws3.row_dimensions[T3_R].height = 26

for i, name in enumerate(C3, start=1):
    c = ws3.cell(H3_R, i, name)
    c.font = Font(bold=True, color=XL_HEADER_FONT)
    c.fill = PatternFill("solid", fgColor=XL_HEADER)
    c.alignment = CENTER
    c.border = BORDER

name_to_path = {name: path for _, path, name, _, _ in MODULES}


def resolve_modules(text: str) -> str:
    """把「中文名 / 中文名」翻成「中文名（路径）」，逐行一条。"""
    lines = []
    for part in text.split(" / "):
        part = part.strip()
        path = name_to_path.get(part)
        lines.append(f"{part}（{path}）" if path else part)
    return "\n".join(lines)


chain_order = []
for ch, *_ in CHAINS:
    if ch not in chain_order:
        chain_order.append(ch)

for i, (chain, step, what, mods) in enumerate(CHAINS):
    r = D3_R0 + i
    ws3.cell(r, 1, i + 1)
    ws3.cell(r, 2, chain)
    ws3.cell(r, 3, step)
    ws3.cell(r, 4, what)
    ws3.cell(r, 5, mods)
    ws3.cell(r, 6, resolve_modules(mods))
    band = chain_order.index(chain) % 2 == 1
    for col in range(1, len(C3) + 1):
        cell = ws3.cell(r, col)
        cell.border = BORDER
        cell.alignment = CENTER if col in (1, 3) else WRAP_L
        if col == 6:
            cell.font = Font(name="Consolas", size=9)
        if band:
            cell.fill = PatternFill("solid", fgColor=XL_BAND)

LAST3 = D3_R0 + len(CHAINS) - 1
ws3.auto_filter.ref = f"A{H3_R}:F{LAST3}"
ws3.freeze_panes = f"A{D3_R0}"
for col, w in zip("ABCDEF", [6, 26, 8, 56, 30, 46]):
    ws3.column_dimensions[col].width = w

wb.properties.title = "鬼畜哈吉米 · 模块地图"
wb.save(OUT)
print("OK ->", os.path.abspath(OUT))
print("sheets:", wb.sheetnames)
print("模块地图行数:", len(MODULES), "| 调用链行数:", len(CHAINS))
